#!/usr/bin/env python3
"""Build and style the collected Aletheia Word rendition."""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Aletheia_Protocol_0.5_RC2_Complete_Collected_Documents.md"
OUTPUT = ROOT / "Aletheia_Protocol_0.5_RC2_Complete_Collected_Documents.docx"

INK = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x66, 0x66)


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_style(style, *, font="Calibri", size=11, color=None, bold=None,
                    before=0, after=6, line=1.25, keep_next=False):
    style.font.name = font
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next


def get_style(styles, display_name: str, style_id: str | None = None):
    for style in styles:
        if style.name == display_name or (style_id is not None and style.style_id == style_id):
            return style
    raise KeyError(display_name)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_paragraph(paragraph, fill="F4F6F9"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr, fld_char_2])


def word_markdown() -> str:
    text = SOURCE.read_text(encoding="utf-8")
    first_div = text.find('<div id="')
    if first_div < 0:
        raise SystemExit("Collected source has no controlled sections")
    body = text[first_div:]
    sections = re.findall(r'<div id="[^"]+">\s*(.*?)\s*</div>', body, flags=re.S)
    if not sections:
        raise SystemExit("Collected source sections could not be parsed")

    rendered: list[str] = []
    for section in sections:
        lines = section.strip().splitlines()
        yaml_lines: list[str] = []
        if lines and lines[0].strip() == "---":
            close = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), None)
            if close is not None and any(":" in line for line in lines[1:close]):
                yaml_lines = lines[1:close]
                lines = lines[close + 1:]

        if yaml_lines:
            heading = next((i for i, line in enumerate(lines) if line.startswith("# ")), None)
            yaml_block = ["```yaml", *yaml_lines, "```"]
            if heading is not None:
                lines = lines[:heading + 1] + [""] + yaml_block + [""] + lines[heading + 1:]
            else:
                lines = yaml_block + [""] + lines
        rendered.append("\n".join(lines).strip())

    contents = """# Contents

1. Overview and release status
2. Aletheia Principles v1.0
3. Protocol Specification 0.5 RC2
4. Agentic Systems Profile 0.1 RC2
5. Agentic Systems proposal and evidence review
6. Adoption and AI handover guide
7. Governance and contributing
8. Repository and publication policy
9. Licensing and open adoption
10. Changelog and migration
11. Project, node, handover, provenance and action-receipt templates
12. Minimal and agentic-system worked examples
13. All-in-One AI Bootstrap
"""
    return contents + "\n\n" + "\n\n".join(rendered) + "\n"


def build_raw(markdown_path: Path, raw_path: Path):
    subprocess.run(
        [
            "pandoc",
            str(markdown_path),
            "--from=gfm+raw_html",
            "--to=docx",
            "--standalone",
            "--metadata", "title=Aletheia Protocol 0.5 RC2",
            "--metadata", "subtitle=Complete Collected Documents",
            "--metadata", "author=Knowledge belongs to the project, not the Assistant.",
            "--metadata", "date=Release candidate · 11 August 2026 · Originating Steward: Karsten Evans",
            "--output", str(raw_path),
        ],
        check=True,
    )


def style_document(raw_path: Path):
    doc = Document(raw_path)
    doc.core_properties.title = "Aletheia Protocol 0.5 RC2 — Complete Collected Documents"
    doc.core_properties.subject = "Human-and-AI knowledge stewardship protocol"
    doc.core_properties.author = "Karsten Evans with human–AI collaborators"
    doc.core_properties.comments = "Licensed release candidate; independent review pending"

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True

        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Aletheia Protocol 0.5 RC2  ·  Complete Collected Documents"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        for run in hp.runs:
            set_font(run, size=8.5, color=MUTED)

        footer = section.footer
        fp = footer.paragraphs[0]
        add_page_number(fp)

    styles = doc.styles
    normal_style = get_style(styles, "Normal")
    title_style = get_style(styles, "Title")
    h1_style = get_style(styles, "Heading 1", "Heading1")
    h2_style = get_style(styles, "Heading 2", "Heading2")
    h3_style = get_style(styles, "Heading 3", "Heading3")
    configure_style(normal_style, size=11, after=6, line=1.25)
    configure_style(title_style, size=30, color=INK, bold=True, after=10, line=1.0)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(128)
    if "Subtitle" in styles:
        configure_style(styles["Subtitle"], size=16, color=BLUE, after=22, line=1.0)
        styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Author" in styles:
        configure_style(styles["Author"], size=11, color=DARK_BLUE, after=8, line=1.15)
        styles["Author"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "Date" in styles:
        configure_style(styles["Date"], size=10.5, color=MUTED, after=0, line=1.15)
        styles["Date"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    configure_style(h1_style, size=16, color=BLUE, bold=True, before=18, after=10, line=1.0, keep_next=True)
    configure_style(h2_style, size=13, color=BLUE, bold=True, before=14, after=7, line=1.0, keep_next=True)
    configure_style(h3_style, size=12, color=DARK_BLUE, bold=True, before=10, after=5, line=1.0, keep_next=True)
    if "Heading 4" in styles:
        configure_style(get_style(styles, "Heading 4", "Heading4"), size=11, color=INK, bold=True, before=8, after=4, line=1.0, keep_next=True)
    if "TOC Heading" in styles:
        configure_style(get_style(styles, "TOC Heading", "TOCHeading"), size=16, color=BLUE, bold=True, before=0, after=10, line=1.0, keep_next=True)

    if "Source Code" in styles:
        configure_style(styles["Source Code"], font="Consolas", size=8.5, color=INK, after=3, line=1.0)
    if "Verbatim Char" in styles:
        verbatim = styles["Verbatim Char"]
        verbatim.font.name = "Consolas"
        verbatim._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
        verbatim._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        verbatim.font.size = Pt(8.5)
        verbatim.font.color.rgb = INK
    if "Block Text" in styles:
        configure_style(styles["Block Text"], size=10.5, color=DARK_BLUE, after=8, line=1.15)
        styles["Block Text"].paragraph_format.left_indent = Inches(0.25)

    section_break_titles = {
        "Contents",
        "Aletheia Protocol 0.5 RC2 — Public Review Candidate",
        "Aletheia Principles v1.0",
        "Aletheia Protocol Specification 0.5 RC2",
        "Aletheia Agentic Systems Profile 0.1 RC2",
        "PROP-AGENTIC-001 — Agentic Systems and Population-Level Provenance",
        "Adoption and AI Handover Guide",
        "Governance and Contributing",
        "Repository and Publication Policy",
        "Licensing",
        "Changelog",
        "Project Specification",
        "Minimal Worked Example",
        "Aletheia Protocol — All-in-One AI Bootstrap",
    }
    date_break_added = False
    in_principles = False
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name == "Date" and not date_break_added:
            paragraph.add_run().add_break()
            paragraph.runs[-1]._r[-1].set(qn("w:type"), "page")
            date_break_added = True
        if style_name == "Heading 1":
            in_principles = paragraph.text.strip() == "Aletheia Principles v1.0"
            paragraph.paragraph_format.page_break_before = paragraph.text.strip() in section_break_titles
            paragraph.paragraph_format.keep_with_next = True
        elif style_name.startswith("Heading") or style_name == "TOC Heading":
            paragraph.paragraph_format.keep_with_next = True

        p_pr = paragraph._p.pPr
        num_pr = p_pr.numPr if p_pr is not None else None
        if num_pr is not None:
            level = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
            paragraph.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
            paragraph.paragraph_format.first_line_indent = Inches(-0.188)
            paragraph.paragraph_format.space_after = Pt(2 if in_principles else 4)
            paragraph.paragraph_format.line_spacing = 1.12 if in_principles else 1.25
            if in_principles:
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)

        if style_name == "Source Code":
            shade_paragraph(paragraph)
            paragraph.paragraph_format.left_indent = Inches(0.12)
            paragraph.paragraph_format.right_indent = Inches(0.12)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True

    for table in doc.tables:
        table.autofit = False
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                set_cell_margins(cell)
                cell.vertical_alignment = 1
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    if row_idx == 0:
                        for run in p.runs:
                            run.bold = True

    doc.save(OUTPUT)


def audit(docx_path: Path):
    doc = Document(docx_path)
    section = doc.sections[0]
    assert round(section.page_width.inches, 3) == 8.5
    assert round(section.page_height.inches, 3) == 11.0
    assert all(round(value.inches, 3) == 1.0 for value in (
        section.top_margin, section.right_margin, section.bottom_margin, section.left_margin
    ))
    assert get_style(doc.styles, "Normal").font.name == "Calibri"
    assert get_style(doc.styles, "Normal").font.size.pt == 11
    assert get_style(doc.styles, "Heading 1", "Heading1").font.size.pt == 16
    assert get_style(doc.styles, "Heading 2", "Heading2").font.size.pt == 13
    assert get_style(doc.styles, "Heading 3", "Heading3").font.size.pt == 12
    assert any(p.style and p.style.name == "Title" for p in doc.paragraphs)
    assert any(p.style and p.style.name == "Heading 1" for p in doc.paragraphs)
    print(f"Preset audit passed: compact_reference_guide + editorial_cover override ({docx_path})")


if __name__ == "__main__":
    with tempfile.TemporaryDirectory(prefix="aletheia-docx-") as tmp:
        tmp_path = Path(tmp)
        md_path = tmp_path / "word_source.md"
        raw_path = tmp_path / "raw.docx"
        md_path.write_text(word_markdown(), encoding="utf-8")
        build_raw(md_path, raw_path)
        style_document(raw_path)
    audit(OUTPUT)
    print(OUTPUT)
